import io

import pdfplumber
import requests
from docx import Document

from text_cleaner import clean_text


def download_file(url: str) -> bytes:
    """Download a Cloudinary file into an in-memory buffer."""
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    return response.content


def extract_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from PDF pages and tables using pdfplumber.
    """
    text_parts = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    row_text = ' | '.join(
                        cell.strip() for cell in row if cell and cell.strip()
                    )
                    if row_text:
                        text_parts.append(row_text)

    return clean_text('\n'.join(text_parts))


def extract_from_docx(file_bytes: bytes) -> str:
    """
    Extract text from DOCX paragraphs and tables using python-docx.
    """
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = ' | '.join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                text_parts.append(row_text)

    return clean_text('\n'.join(text_parts))


def extract_text(file_url: str, file_type: str) -> str:
    """
    Download a resume from a URL and route it to the correct extractor.
    """
    file_bytes = download_file(file_url)

    if file_type == 'pdf':
        return extract_from_pdf(file_bytes)
    if file_type == 'docx':
        return extract_from_docx(file_bytes)

    raise ValueError(f'Unsupported file type: {file_type}')
